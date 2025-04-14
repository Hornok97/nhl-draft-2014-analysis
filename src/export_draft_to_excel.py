import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

# 1️⃣ Načti CSV
df = pd.read_csv("nhl_draft_2014.csv")

# 2️⃣ Vytvoř různé pohledy
df_top_points = df.copy()
df_top_points["Points"] = pd.to_numeric(df_top_points["Points"], errors="coerce")
df_top_points = df_top_points.sort_values(by="Points", ascending=False)

df_top_gp = df.copy()
df_top_gp["GP"] = pd.to_numeric(df_top_gp["GP"], errors="coerce")
df_top_gp = df_top_gp.sort_values(by="GP", ascending=False)

df_defenders = df[df["Player"].str.contains(r"\(D\)", na=False)].copy()
df_defenders["Points"] = pd.to_numeric(df_defenders["Points"], errors="coerce")
df_defenders = df_defenders.sort_values(by="Points", ascending=False)

# 3️⃣ Nahraď NaN za '-'
tables = {
    "Draft2014": df,
    "Top podle bodů": df_top_points,
    "Počet zápasů": df_top_gp,
    "Obránci podle bodů": df_defenders
}

for name in tables:
    tables[name] = tables[name].fillna("-")

# 4️⃣ Ulož do Excelu
filename = "nhl_draft_2014_analýza.xlsx"
with pd.ExcelWriter(filename, engine="openpyxl") as writer:
    for sheet, data in tables.items():
        data.to_excel(writer, sheet_name=sheet, index=False)

# 5️⃣ Načti workbook a zarovnej styl
wb = load_workbook(filename)

for sheet in wb.sheetnames:
    ws = wb[sheet]
    # Zarovnání všech buněk
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
    # Tučný záhlaví
    for cell in ws[1]:
        cell.font = Font(bold=True)

wb.save(filename)
print(f"Excel '{filename}' byl úspěšně vytvořen a naformátován.")
